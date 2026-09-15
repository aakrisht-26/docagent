"""
ExcelReaderSkill — extracts structured text and data from Excel and CSV files.

Supported formats:
    .xlsx          — via openpyxl (structure-aware, handles merged cells)
    .xls           — via xlrd (legacy BIFF8; values only)
    .csv           — via pandas (encoding auto-detection)

Each sheet becomes one DocumentChunk. Tables are preserved as text representations
for downstream summarization and question extraction.
"""

from __future__ import annotations

import io
import time
from pathlib import Path
from typing import Any, Dict, List, Optional

from core.models import DocumentChunk, ParsedDocument, SkillInput, SkillOutput
from skills.base_skill import BaseSkill
from utils.logger import get_logger

logger = get_logger(__name__)

#: The first bytes of the two workbook containers. The engine is chosen from
#: these, not from the file's name: openpyxl refuses by extension before it reads
#: a byte, and spreadsheets in the wild are often mislabelled.
_OLE2_MAGIC = bytes.fromhex("d0cf11e0a1b11ae1")   # legacy .xls: BIFF8 in an OLE2 compound file
_ZIP_MAGIC = b"PK\x03\x04"                         # .xlsx


class UnreadableWorkbook(ValueError):
    """A spreadsheet that cannot be read, already described for the user."""


def _xls_cell_text(cell, datemode: int) -> str:
    """One xlrd cell as the text the .xlsx path produces for the same cell.

    xlrd hands back every number as a float and every date as a serial. Copied
    as they came, 400000 read "400000.0" and a date read "46082.0", so one
    workbook read differently as .xls and as .xlsx. Measured identical after
    this on the committed samples, a 30,000-row workbook, and one holding dates,
    times of day, booleans, a formula error and a merged cell.
    """
    import xlrd

    kind = cell.ctype
    if kind in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return ""
    if kind == xlrd.XL_CELL_NUMBER:
        return str(int(cell.value)) if float(cell.value).is_integer() else str(cell.value)
    if kind == xlrd.XL_CELL_DATE:
        moment = xlrd.xldate.xldate_as_datetime(cell.value, datemode)
        # Under 1 is a time of day with no date, which openpyxl returns as a time.
        return str(moment.time()) if cell.value < 1 else str(moment)
    if kind == xlrd.XL_CELL_BOOLEAN:
        return str(bool(cell.value))
    if kind == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(cell.value, "#ERROR")
    return str(cell.value).strip()


def _not_a_workbook(name: str, head: bytes) -> str:
    start = head.lstrip(b"\xef\xbb\xbf \t\r\n").lower()
    if start.startswith((b"<html", b"<!doctype", b"<?xml", b"<table")):
        what = "a web page or XML file saved with a spreadsheet name"
    else:
        what = "not an Excel workbook"
    return (f"'{name}' has an Excel file name but is {what}, so it cannot be read "
            f"as a spreadsheet. Open it in Excel and save it as .xlsx, or export it as CSV.")


def _damaged_xls(name: str, exc: Exception) -> str:
    if "encrypted" in str(exc).lower():
        return (f"'{name}' is a password-protected .xls workbook and cannot be read. "
                f"Remove the password in Excel and upload it again.")
    return (f"'{name}' is a legacy .xls workbook that could not be read; it may be "
            f"damaged. Opening it in Excel and saving it as .xlsx usually recovers it. "
            f"(detail: {type(exc).__name__}: {exc})")


class ExcelReaderSkill(BaseSkill):
    """
    Reads and parses Excel / CSV files into a structured ParsedDocument.

    Config keys:
        max_sheets        (int)  : Max number of sheets to process (default: 50)
        include_formulas  (bool) : Include formula text instead of values (default: False)
        max_rows_per_sheet(int)  : Row limit per sheet (default: 10000)
    """

    name = "excel_reader"
    description = "Extracts data, headers, and text from Excel / CSV files, per sheet."
    required_inputs = ["file_path"]

    def __init__(self, config: Optional[Dict[str, Any]] = None) -> None:
        super().__init__(config)
        self._max_sheets   = self.get_config("max_sheets", 50)
        self._max_rows     = self.get_config("max_rows_per_sheet", 10_000)
        self._incl_formula = self.get_config("include_formulas", False)

    # ── Skill entry point ─────────────────────────────────────────────

    def execute(self, inputs: SkillInput) -> SkillOutput:
        start = time.monotonic()
        file_path = Path(inputs.data["file_path"])

        if not file_path.exists():
            return SkillOutput(success=False, data=None,
                               error=f"File not found: {file_path}")

        suffix = file_path.suffix.lower()
        try:
            if suffix == ".csv":
                doc = self._parse_csv(file_path)
            else:
                doc = self._parse_excel(file_path)
        except UnreadableWorkbook as exc:
            # Already a sentence written for the user; a traceback helps nobody.
            self.logger.warning(f"Excel read refused: {exc}")
            return SkillOutput(
                success=False, data=None, error=str(exc),
                duration_ms=(time.monotonic() - start) * 1000,
            )
        except Exception as exc:
            self.logger.error(f"Excel read failed: {exc}", exc_info=True)
            return SkillOutput(
                success=False, data=None,
                error=f"Failed to parse '{file_path.name}': {exc}",
                duration_ms=(time.monotonic() - start) * 1000,
            )

        warnings: List[str] = []
        if doc.is_empty:
            warnings.append("No data extracted. The file may be empty or all sheets are blank.")

        self.logger.info(
            f"Excel parsed: {doc.page_count} sheet(s), "
            f"{doc.word_count} words, {len(doc.tables)} table(s)"
        )
        return SkillOutput(
            success=True,
            data=doc,
            warnings=warnings,
            duration_ms=(time.monotonic() - start) * 1000,
        )

    # ── Parsers ───────────────────────────────────────────────────────

    def _parse_excel(self, file_path: Path) -> ParsedDocument:
        """Parse a workbook, choosing the engine from its bytes, not its name.

        .xls used to go to openpyxl with everything else, and openpyxl cannot
        read the legacy BIFF8 format at all, so every real .xls upload failed.

            OLE2 header (d0 cf 11 e0 ...)  -> xlrd      legacy .xls
            ZIP header (PK 03 04)          -> openpyxl  .xlsx, whatever its name
            anything else                  -> refused, in words a user can act on

        Both engines produce the same rows, so a workbook reads the same saved
        either way.
        """
        with open(file_path, "rb") as fh:
            head = fh.read(16)
        if head[:8] == _OLE2_MAGIC:
            sheets = self._xls_sheets(file_path)
        elif head[:4] == _ZIP_MAGIC:
            sheets = self._xlsx_sheets(file_path)
        else:
            raise UnreadableWorkbook(_not_a_workbook(file_path.name, head))

        chunks: List[DocumentChunk] = []
        tables: List[Dict[str, Any]] = []
        text_parts: List[str] = []
        sheet_names: List[str] = []

        for sheet_idx, (sheet_name, rows) in enumerate(sheets):
            sheet_names.append(sheet_name)

            if not rows:
                self.logger.debug(f"Sheet '{sheet_name}' is empty — skipped.")
                continue

            df = self._rows_to_dataframe(rows)
            text_repr = self._dataframe_to_text(df, sheet_name)
            text_parts.append(text_repr)

            tables.append({
                "sheet": sheet_name,
                "sheet_index": sheet_idx,
                "rows": len(df),
                "columns": len(df.columns),
                "preview": rows[:5],
            })
            chunks.append(DocumentChunk(
                text=text_repr,
                page_or_sheet=sheet_name,
                chunk_index=sheet_idx,
                metadata={"sheet": sheet_name, "rows": len(rows)},
            ))

        return ParsedDocument(
            file_name=file_path.name,
            file_type="excel",
            chunks=chunks,
            full_text="\n\n" + ("─" * 60) + "\n\n".join(text_parts),
            tables=tables,
            metadata={"sheets": len(sheet_names), "sheet_names": sheet_names},
            page_count=len(sheet_names) if sheet_names else 0,
            sheet_names=sheet_names,
        )

    def _xlsx_sheets(self, file_path: Path):
        """(sheet name, non-empty rows) for each sheet, via openpyxl.

        Opened from a file object, not the path: openpyxl refuses by EXTENSION
        before it reads a byte, so an .xlsx saved or renamed as .xls failed on
        its name alone.
        """
        import openpyxl

        with open(file_path, "rb") as fh:
            wb = openpyxl.load_workbook(fh, data_only=not self._incl_formula)
            try:
                for sheet_name in wb.sheetnames[: self._max_sheets]:
                    rows: List[List[str]] = []
                    for row in wb[sheet_name].iter_rows(max_row=self._max_rows, values_only=True):
                        if any(cell is not None for cell in row):
                            rows.append([str(c).strip() if c is not None else "" for c in row])
                    yield sheet_name, rows
            finally:
                wb.close()

    def _xls_sheets(self, file_path: Path):
        """(sheet name, non-empty rows) for each sheet of a legacy .xls, via xlrd.

        Values only: xlrd has no formula text for BIFF8, so `include_formulas`
        cannot apply, and cached values are what the .xlsx path returns by
        default.

        xlrd's failures are not passed on as they come. A truncated file raised
        a bare "IndexError: array index out of range" after printing OLE2
        warnings to stdout; it now becomes one sentence, and the warnings go to
        a buffer.
        """
        import xlrd

        try:
            book = xlrd.open_workbook(str(file_path), on_demand=True, logfile=io.StringIO())
        except Exception as exc:
            raise UnreadableWorkbook(_damaged_xls(file_path.name, exc)) from exc
        try:
            for sheet_name in book.sheet_names()[: self._max_sheets]:
                try:
                    sheet = book.sheet_by_name(sheet_name)
                except Exception as exc:
                    raise UnreadableWorkbook(_damaged_xls(file_path.name, exc)) from exc
                rows = []
                for r in range(min(sheet.nrows, self._max_rows)):
                    values = [_xls_cell_text(cell, book.datemode) for cell in sheet.row(r)]
                    if any(values):
                        rows.append(values)
                book.unload_sheet(sheet_name)
                yield sheet_name, rows
        finally:
            book.release_resources()

    def _parse_csv(self, file_path: Path) -> ParsedDocument:
        """Parse .csv files using pandas with encoding fallback."""
        import pandas as pd

        for enc in ("utf-8", "latin-1", "cp1252"):
            try:
                df = pd.read_csv(str(file_path), nrows=self._max_rows, encoding=enc)
                break
            except (UnicodeDecodeError, Exception):
                continue
        else:
            raise ValueError("Could not decode CSV with any known encoding.")

        text_repr = self._dataframe_to_text(df, file_path.stem)
        return ParsedDocument(
            file_name=file_path.name,
            file_type="excel",
            chunks=[DocumentChunk(
                text=text_repr,
                page_or_sheet=1,
                chunk_index=0,
                metadata={"source": "csv", "rows": len(df), "cols": len(df.columns)},
            )],
            full_text=text_repr,
            tables=[{"sheet": "CSV", "rows": len(df), "columns": len(df.columns)}],
            metadata={"format": "csv", "rows": len(df), "columns": list(df.columns)},
            page_count=1,
            sheet_names=["CSV"],
        )

    # ── Helpers ───────────────────────────────────────────────────────

    @staticmethod
    def _rows_to_dataframe(rows: List[List[str]]):
        """Convert list-of-rows to a DataFrame, promoting first row to header if suitable."""
        import pandas as pd

        df = pd.DataFrame(rows)
        if rows and ExcelReaderSkill._looks_like_header(rows[0]):
            df.columns = [str(h).strip() or f"col_{i}" for i, h in enumerate(rows[0])]
            df = df.iloc[1:].reset_index(drop=True)
        return df

    @staticmethod
    def _looks_like_header(row: List[str]) -> bool:
        """Heuristic: header rows have mostly non-numeric text."""
        non_empty = [c for c in row if c.strip()]
        if not non_empty:
            return False
        numeric = sum(
            1 for c in non_empty
            if c.replace(".", "").replace("-", "").replace(",", "").isdigit()
        )
        return numeric / len(non_empty) < 0.5

    @staticmethod
    def _dataframe_to_text(df: pd.DataFrame, sheet_name: str) -> str:
        """Convert DataFrame to a human-readable text block with smart sampling."""
        import pandas as pd

        row_count = len(df)
        lines = [f"[Sheet: {sheet_name}]", f"- Total Rows: {row_count}", ""]
        
        try:
            # If the dataset is large, take head + tail to show breadth
            if row_count > 600:
                head = df.head(300)
                tail = df.tail(300)
                lines.append(head.to_string(index=False, max_cols=30))
                lines.append(f"\n... [Rows 301 to {row_count-300} omitted] ...\n")
                lines.append(tail.to_string(index=False, max_cols=30))
            else:
                lines.append(df.to_string(index=False, max_cols=30))
        except Exception:
            lines.append(str(df.head(100)))
            
        return "\n".join(lines)
